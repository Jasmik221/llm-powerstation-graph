from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\demon\Documents\ChatGPT\Inżynierka")
OUT = ROOT / "materiały_na_spotkanie"
ASSETS = OUT / "assets"
OUT.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)

DOCX_PATH = OUT / "Raport_poczatkowy_pracy_inzynierskiej.docx"

NAVY = "16324F"
BLUE = "2E75B6"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F3F5F7"
MID_GRAY = "66727C"
LIGHT_BORDER = "D9D9D9"
GREEN = "2E8B67"
ORANGE = "D9822B"
RED = "C34A4A"
WHITE = "FFFFFF"
BLACK = "000000"


def font(size: int, bold: bool = False):
    path = r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"
    return ImageFont.truetype(path, size)


def rounded_box(draw, xy, fill, outline, radius=22, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def centered_text(draw, xy, text, fnt, fill="#102A3A", line_gap=8):
    x1, y1, x2, y2 = xy
    lines = text.split("\n")
    boxes = [draw.textbbox((0, 0), line, font=fnt) for line in lines]
    heights = [b[3] - b[1] for b in boxes]
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + (y2 - y1 - total_h) / 2
    for line, b, h in zip(lines, boxes, heights):
        w = b[2] - b[0]
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=fnt, fill=fill)
        y += h + line_gap


def arrow(draw, start, end, color="#2E75B6", width=5):
    draw.line([start, end], fill=color, width=width)
    import math
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 18
    a1 = angle + math.pi * 0.82
    a2 = angle - math.pi * 0.82
    p1 = (end[0] + size * math.cos(a1), end[1] + size * math.sin(a1))
    p2 = (end[0] + size * math.cos(a2), end[1] + size * math.sin(a2))
    draw.polygon([end, p1, p2], fill=color)


def make_architecture(path: Path):
    img = Image.new("RGB", (1800, 1030), "#F7F9FB")
    d = ImageDraw.Draw(img)
    title = font(42, True)
    label = font(28, True)
    body = font(22)
    d.text((80, 45), "Architektura prototypu", font=title, fill="#000000")
    d.text((80, 105), "Dwie niezależne ścieżki tworzą grafy możliwe do bezpośredniego porównania", font=body, fill="#66727C")

    top = [
        (80, 230, 330, 365, "pandapower\ncase14()", "#EAF2F8", "#2E75B6"),
        (420, 230, 670, 365, "Tabele\nDataFrame", "#FFFFFF", "#AAB7C2"),
        (760, 230, 1010, 365, "Konwerter\nPython", "#DFF6F5", "#1B9AAA"),
        (1100, 230, 1370, 365, "Graf\nreferencyjny", "#E5F4EC", "#2E8B67"),
    ]
    bottom = [
        (80, 610, 330, 745, "Opis\ntekstowy", "#FFF2DF", "#D9822B"),
        (420, 610, 670, 745, "Adapter LLM\nAPI lub Ollama", "#FFFFFF", "#AAB7C2"),
        (760, 610, 1010, 745, "JSON i\nwalidacja", "#DFF6F5", "#1B9AAA"),
        (1100, 610, 1370, 745, "Graf\nwygenerowany", "#E5F4EC", "#2E8B67"),
    ]
    for row in (top, bottom):
        for x1, y1, x2, y2, text, fill, outline in row:
            rounded_box(d, (x1, y1, x2, y2), fill, outline)
            centered_text(d, (x1, y1, x2, y2), text, label)
        for a, b in zip(row[:-1], row[1:]):
            arrow(d, (a[2] + 10, (a[1] + a[3]) // 2), (b[0] - 10, (b[1] + b[3]) // 2), "#2E75B6")

    rounded_box(d, (1490, 390, 1725, 585), "#16324F", "#16324F", radius=95)
    centered_text(d, (1490, 390, 1725, 585), "Neo4j", font(36, True), "#FFFFFF")
    arrow(d, (1380, 300), (1500, 430), "#2E8B67")
    arrow(d, (1380, 678), (1500, 548), "#2E8B67")
    d.text((1125, 835), "Porównanie grafów", font=label, fill="#16324F")
    d.text((1125, 880), "metryki, wizualizacja i analiza błędów", font=body, fill="#66727C")
    img.save(path, quality=95)


def make_graph_schema(path: Path):
    img = Image.new("RGB", (1800, 1150), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((70, 35), "Wstępny model grafu właściwościowego", font=font(42, True), fill="#000000")
    node_font = font(25, True)
    small = font(18)
    rel = font(17, True)

    boxes = {
        "gen": (70, 235, 400, 370, "Generator lub ExtGrid\nsource_id, pp_index, p_mw, vm_pu", "#E5F4EC", "#2E8B67"),
        "load": (70, 520, 400, 655, "Load\nsource_id, pp_index, p_mw, q_mvar", "#FFF2DF", "#D9822B"),
        "bus1": (640, 180, 1000, 325, "Bus\nsource_id, pp_index, name, vn_kv, geo", "#EAF2F8", "#2E75B6"),
        "line": (640, 445, 1000, 610, "Line\nsource_id, pp_index, length_km, r, x, max_i_ka", "#DFF6F5", "#1B9AAA"),
        "trafo": (640, 790, 1000, 975, "Transformer\nsource_id, pp_index, sn_mva, napięcia, vk, tap", "#FBE8E8", "#C34A4A"),
        "bus2": (1280, 180, 1640, 325, "Bus\nsource_id, pp_index, name, vn_kv, geo", "#EAF2F8", "#2E75B6"),
        "shunt": (1280, 520, 1640, 655, "Shunt\nsource_id, pp_index, p_mw, q_mvar, step", "#F3F0FA", "#7356A8"),
    }
    for key, (x1, y1, x2, y2, text, fill, outline) in boxes.items():
        rounded_box(d, (x1, y1, x2, y2), fill, outline)
        lines = text.split("\n")
        centered_text(d, (x1, y1 + 8, x2, y1 + 65), lines[0], node_font)
        centered_text(d, (x1 + 12, y1 + 62, x2 - 12, y2 - 8), "\n".join(lines[1:]), small, "#425563", 4)

    arrow(d, (410, 300), (630, 255), "#2E8B67")
    d.text((430, 235), "CONNECTED_TO", font=rel, fill="#2E8B67")
    arrow(d, (410, 585), (630, 300), "#D9822B")
    d.text((430, 490), "CONNECTED_TO", font=rel, fill="#D9822B")
    arrow(d, (790, 435), (820, 335), "#1B9AAA")
    d.text((635, 365), "FROM_BUS", font=rel, fill="#1B9AAA")
    arrow(d, (1010, 500), (1270, 290), "#1B9AAA")
    d.text((1060, 390), "TO_BUS", font=rel, fill="#1B9AAA")
    arrow(d, (780, 780), (820, 335), "#C34A4A")
    d.text((655, 670), "HV_BUS", font=rel, fill="#C34A4A")
    arrow(d, (1010, 850), (1270, 325), "#C34A4A")
    d.text((1080, 690), "LV_BUS", font=rel, fill="#C34A4A")
    arrow(d, (1270, 585), (1010, 295), "#7356A8")
    d.text((1080, 470), "CONNECTED_TO", font=rel, fill="#7356A8")
    d.text((70, 1055), "Identyfikator techniczny: source_id = <sieć>:<typ>:<pp_index>, np. case14:bus:0", font=font(22, True), fill="#16324F")
    img.save(path, quality=95)


def make_ui_mock(path: Path):
    img = Image.new("RGB", (1800, 1080), "#EEF2F5")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((65, 55, 1735, 1025), radius=24, fill="#FFFFFF", outline="#CAD4DC", width=3)
    d.rectangle((65, 55, 1735, 145), fill="#16324F")
    d.text((105, 83), "EnergyGraph LLM", font=font(31, True), fill="#FFFFFF")
    d.rounded_rectangle((1430, 79, 1685, 124), radius=20, fill="#E5F4EC")
    centered_text(d, (1430, 79, 1685, 124), "Neo4j połączony", font(18, True), "#2E8B67")

    d.text((105, 185), "Opis sieci", font=font(25, True), fill="#16324F")
    d.rounded_rectangle((100, 235, 590, 600), radius=16, fill="#F8FAFB", outline="#CAD4DC", width=2)
    text = ["Sieć zawiera węzły B1, B2 i B3.", "Węzły B1 i B2 łączy linia", "o długości 10 km. Do B3", "przyłączony jest odbiór 20 MW."]
    y = 270
    for line in text:
        d.text((130, y), line, font=font(21), fill="#526471")
        y += 43
    d.rounded_rectangle((100, 630, 330, 690), radius=18, fill="#2E75B6")
    centered_text(d, (100, 630, 330, 690), "Generuj graf", font(21, True), "#FFFFFF")
    d.rounded_rectangle((350, 630, 590, 690), radius=18, fill="#E5F4EC", outline="#2E8B67")
    centered_text(d, (350, 630, 590, 690), "Walidacja poprawna", font(18, True), "#2E8B67")
    d.text((105, 735), "Model", font=font(19, True), fill="#66727C")
    d.rounded_rectangle((100, 775, 590, 840), radius=14, fill="#FFFFFF", outline="#CAD4DC", width=2)
    d.text((130, 794), "Ollama / Qwen", font=font(20), fill="#16324F")

    d.text((650, 185), "Podgląd grafu", font=font(25, True), fill="#16324F")
    d.rounded_rectangle((640, 235, 1320, 900), radius=16, fill="#F8FAFB", outline="#CAD4DC", width=2)
    nodes = {
        "B1": (745, 370, "#EAF2F8", "#2E75B6"),
        "L1": (940, 315, "#DFF6F5", "#1B9AAA"),
        "B2": (1130, 370, "#EAF2F8", "#2E75B6"),
        "T1": (1080, 575, "#FBE8E8", "#C34A4A"),
        "B3": (970, 760, "#EAF2F8", "#2E75B6"),
        "Load": (760, 720, "#FFF2DF", "#D9822B"),
    }
    links = [("B1", "L1", "#1B9AAA"), ("L1", "B2", "#1B9AAA"), ("B2", "T1", "#C34A4A"), ("T1", "B3", "#C34A4A"), ("Load", "B3", "#D9822B")]
    for a, b, color in links:
        pa, pb = nodes[a], nodes[b]
        arrow(d, (pa[0], pa[1]), (pb[0], pb[1]), color, 4)
    for name, (x, y, fill, outline) in nodes.items():
        r = 55 if name != "Load" else 63
        d.ellipse((x-r, y-r, x+r, y+r), fill=fill, outline=outline, width=4)
        centered_text(d, (x-r, y-r, x+r, y+r), name, font(19, True), outline)

    d.text((1380, 185), "Szczegóły elementu", font=font(25, True), fill="#16324F")
    d.text((1380, 255), "Typ: Line\n\nID: case14:line:0\n\nOd: bus 0\n\nDo: bus 1\n\nDługość: 10 km\n\nStatus: poprawny", font=font(20), fill="#526471", spacing=9)
    d.rounded_rectangle((1380, 790, 1655, 850), radius=18, fill="#EAF2F8", outline="#2E75B6")
    centered_text(d, (1380, 790, 1655, 850), "Eksport JSON", font(20, True), "#2E75B6")
    img.save(path, quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Arial", size=10.5, bold=False, color=BLACK, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_table(table, widths=None, center_cols=()):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, NAVY if r_idx == 0 else (PALE_BLUE if r_idx % 2 == 0 else WHITE))
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=9.2 if len(table.columns) >= 4 else 9.6,
                                 bold=r_idx == 0, color=WHITE if r_idx == 0 else BLACK)


def add_table(doc, headers, rows, widths=None, center_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
    style_table(table, widths, center_cols)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.style = doc.styles["Body Text"]
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        r = p.add_run(item)
        set_run_font(r)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run(text)
    set_run_font(r, size=9, italic=True, color=MID_GRAY)
    return p


def add_code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F4F6")
    set_cell_border(cell, "C8D0D7")
    set_cell_margins(cell, 130, 150, 130, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for line in text.split("\n"):
        r = p.add_run(line)
        set_run_font(r, name="Consolas", size=8.5, color="233746")
        r.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Strona ")
    set_run_font(run, size=9, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def build_document():
    architecture = ASSETS / "architektura.png"
    schema = ASSETS / "schemat_grafu.png"
    ui = ASSETS / "makieta_aplikacji.png"
    make_architecture(architecture)
    make_graph_schema(schema)
    make_ui_mock(ui)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(27)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(BLACK)
    title_style.paragraph_format.space_after = Pt(16)

    for style_name, size, before, after in (("Heading 1", 17, 14, 7), ("Heading 2", 13, 10, 5), ("Heading 3", 11.5, 8, 4)):
        st = styles[style_name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(BLACK)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    body = styles["Body Text"]
    body.font.name = "Arial"
    body._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    body._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    body.font.size = Pt(10.5)
    body.font.color.rgb = RGBColor.from_string(BLACK)
    body.paragraph_format.line_spacing = 1.15
    body.paragraph_format.space_after = Pt(6)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_page_number(section.footer.paragraphs[0])

    # Cover page
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(90)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("RAPORT POCZĄTKOWY")
    set_run_font(r, size=12, bold=True, color=BLUE)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("System generowania grafów infrastruktury energetycznej z wykorzystaniem modeli LLM")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("Stan przygotowań i koncepcja prototypu")
    set_run_font(r, size=15, color=MID_GRAY)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(170)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Temat pracy")
    set_run_font(r, size=10, bold=True, color=MID_GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Implementacja modeli LLM wspierających infrastrukturę energetyczną z użyciem baz grafowych")
    set_run_font(r, size=12, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(105)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Materiał na spotkanie z promotorem\n21 października 2026")
    set_run_font(r, size=10.5, color=MID_GRAY)
    doc.add_page_break()

    # Opening summary
    doc.add_heading("Cel dokumentu", level=1)
    add_body(doc, "Dokument podsumowuje decyzje podjęte na początku pracy inżynierskiej oraz przedstawia proponowany kierunek implementacji. Głównym zadaniem systemu będzie przekształcanie tekstowych opisów infrastruktury elektroenergetycznej w ustrukturyzowane grafy przechowywane w Neo4j. Pierwsza wersja prototypu wykorzysta sieć pandapower case14 jako źródło danych referencyjnych.")
    add_body(doc, "Najważniejszym kolejnym krokiem jest zatwierdzenie modelu grafowego i minimalnego zakresu danych. Po tej decyzji można zaimplementować deterministyczny konwerter pandapower do Neo4j, a następnie uruchomić niezależną ścieżkę ekstrakcji informacji z tekstu przez LLM.")

    doc.add_heading("Zakres pracy", level=1)
    add_bullets(doc, [
        "Ocena zdolności modeli językowych do rozpoznawania elementów infrastruktury, ich właściwości i wzajemnych połączeń.",
        "Opracowanie procesu konwersji wyekstrahowanych informacji do grafu właściwościowego.",
        "Implementacja prototypu łączącego Python, modele LLM, walidację danych i Neo4j.",
        "Wizualizacja grafu oraz podstawowa analiza topologii i połączeń.",
        "Porównanie grafów wygenerowanych przez LLM z grafem referencyjnym utworzonym bez udziału modelu.",
    ])

    doc.add_heading("Stan przygotowań", level=1)
    add_table(doc, ["Obszar", "Aktualne ustalenie", "Status"], [
        ["Dane startowe", "pandapower.networks.case14()", "ustalone"],
        ["Baza grafowa", "Neo4j i model property graph", "ustalone"],
        ["Backend", "Python", "ustalone"],
        ["Integracja LLM", "Wymienny adapter dla API oraz modeli lokalnych przez Ollama", "koncepcja"],
        ["Format pośredni", "JSON zgodny ze schematem i walidowany przez Pydantic", "koncepcja"],
        ["Ewaluacja", "Porównanie z grafem referencyjnym oraz analiza błędów", "koncepcja"],
        ["Implementacja", "Kod prototypu i eksperymenty", "do wykonania"],
    ], widths=[1.35, 4.35, 1.15], center_cols=(2,))

    doc.add_page_break()
    doc.add_heading("Dane wejściowe pandapower case14", level=1)
    add_body(doc, "Sieć case14 stanowi niewielki, dobrze kontrolowany przykład, który pozwala sprawdzić pełny proces bez wprowadzania na początku dużego zbioru danych. Dane są przechowywane w tabelach pandas DataFrame. Każdy rekord może zostać odwzorowany na węzeł grafu, a kolumny wskazujące szyny umożliwiają utworzenie relacji.")
    add_table(doc, ["Tabela", "Potwierdzona liczba rekordów", "Proponowany typ węzła", "Główne odwołania"], [
        ["net.bus", "14", "Bus", "indeks 0–13, name 1–14, geo"],
        ["net.line", "15", "Line", "from_bus, to_bus"],
        ["net.trafo", "5", "Transformer", "hv_bus, lv_bus"],
        ["net.gen", "4", "Generator", "bus"],
        ["net.load", "11", "Load", "bus"],
        ["net.ext_grid", "do potwierdzenia", "ExternalGrid", "bus"],
        ["net.shunt", "do potwierdzenia", "Shunt", "bus"],
    ], widths=[1.2, 1.25, 1.65, 2.75], center_cols=(1,))
    add_body(doc, "Indeks rekordu i pole name mają inne znaczenie. Dla szyn indeks DataFrame przyjmuje wartości od 0 do 13, natomiast name zawiera oznaczenia od 1 do 14. Relacje trzeba tworzyć na podstawie indeksów używanych w kolumnach from_bus, to_bus, hv_bus, lv_bus i bus. Nazwa powinna pozostać właściwością opisową.", bold_lead="Indeks rekordu i pole name mają inne znaczenie.")
    add_body(doc, "Przed implementacją należy przypiąć wersję pandapower w pliku zależności. Pozwoli to zachować powtarzalność wyników i uniknąć różnic w strukturze danych między wersjami biblioteki.")

    doc.add_heading("Minimalny zakres pierwszej wersji", level=2)
    add_bullets(doc, [
        "Wczytanie tabel bus, line, trafo, gen i load oraz sprawdzenie ext_grid i shunt.",
        "Zachowanie oryginalnego indeksu pandapower w polu pp_index.",
        "Utworzenie stabilnego identyfikatora source_id zawierającego nazwę sieci, typ elementu i indeks.",
        "Przeniesienie parametrów wymaganych do identyfikacji, wizualizacji i pierwszych eksperymentów.",
        "Oddzielenie danych wejściowych od wyników obliczeń rozpływu mocy, takich jak tabele res_bus lub res_line.",
    ])

    doc.add_page_break()
    doc.add_heading("Architektura systemu", level=1)
    add_body(doc, "Prototyp będzie zawierał dwie ścieżki generowania grafu. Pierwsza przekształci dane pandapower bezpośrednio do Neo4j i utworzy wynik referencyjny. Druga przetworzy opis tekstowy przez model LLM, zapisze wynik w JSON, zweryfikuje go i dopiero wtedy utworzy graf. Obie ścieżki powinny korzystać z tego samego modelu domenowego.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(architecture), width=Inches(6.75))
    add_caption(doc, "Rysunek 1  Proponowana architektura prototypu")

    doc.add_heading("Odpowiedzialność komponentów", level=2)
    add_table(doc, ["Komponent", "Odpowiedzialność"], [
        ["Moduł pandapower", "Wczytuje sieć referencyjną i udostępnia tabele DataFrame."],
        ["Konwerter referencyjny", "Mapuje rekordy tabel na węzły, właściwości i relacje Neo4j."],
        ["Adapter LLM", "Zapewnia wspólny interfejs dla zewnętrznego API i lokalnej Ollamy."],
        ["Warstwa promptów", "Przekazuje reguły ekstrakcji, schemat odpowiedzi i przykłady few shot."],
        ["Walidator", "Sprawdza strukturę JSON, typy, jednostki, identyfikatory oraz poprawność relacji."],
        ["Repozytorium Neo4j", "Tworzy graf wyłącznie na podstawie zweryfikowanych obiektów."],
        ["Moduł ewaluacji", "Porównuje graf wygenerowany z grafem referencyjnym i zapisuje błędy."],
        ["Interfejs", "Udostępnia wprowadzanie opisu, konfigurację modelu, wizualizację i eksport."],
    ], widths=[1.8, 5.05])

    doc.add_page_break()
    doc.add_heading("Proces tworzenia grafu referencyjnego", level=1)
    add_body(doc, "Graf referencyjny powstanie klasycznym skryptem Pythona. LLM nie uczestniczy w tym procesie, dlatego wynik może służyć jako punkt odniesienia podczas ewaluacji.")
    add_bullets(doc, [
        "Wczytanie net = pandapower.networks.case14().",
        "Odczyt wybranych tabel i kolumn oraz normalizacja brakujących wartości.",
        "Utworzenie węzłów Bus przed elementami, które wskazują szyny.",
        "Utworzenie węzłów Line, Transformer, Generator, Load i pozostałych urządzeń.",
        "Utworzenie relacji na podstawie indeksów pandapower.",
        "Sprawdzenie liczby węzłów i relacji oraz zapis manifestu importu.",
    ])
    doc.add_heading("Zasada identyfikacji", level=2)
    add_body(doc, "Każdy element powinien mieć stabilny identyfikator niezależny od pola name. Proponowany format source_id zapobiega kolizjom między typami elementów i między różnymi sieciami testowymi.")
    add_code(doc, "case14:bus:0\ncase14:line:0\ncase14:trafo:0\ncase14:load:0")
    doc.add_heading("Kontrole jakości importu", level=2)
    add_table(doc, ["Kontrola", "Oczekiwany wynik"], [
        ["Liczba szyn", "14 węzłów Bus"],
        ["Liczba linii", "15 węzłów Line"],
        ["Liczba transformatorów", "5 węzłów Transformer"],
        ["Liczba generatorów", "4 węzły Generator"],
        ["Liczba odbiorów", "11 węzłów Load"],
        ["Integralność relacji", "Każda relacja wskazuje istniejący węzeł Bus"],
        ["Unikalność", "Brak powtórzonych source_id w obrębie etykiety"],
    ], widths=[2.3, 4.55])

    doc.add_page_break()
    doc.add_heading("Proces generowania grafu z opisu", level=1)
    add_body(doc, "Model językowy nie powinien bezpośrednio wykonywać zapytań zapisujących dane. Jego zadaniem będzie przygotowanie propozycji struktury w kontrolowanym formacie. Backend pozostanie odpowiedzialny za reguły biznesowe, walidację i transakcję Neo4j.")
    add_table(doc, ["Etap", "Działanie", "Możliwy błąd"], [
        ["1", "Użytkownik wprowadza opis infrastruktury", "Brak danych koniecznych do odtworzenia grafu"],
        ["2", "Prompt przekazuje schemat i reguły ekstrakcji", "Niejednoznaczna instrukcja lub zbyt mało przykładów"],
        ["3", "LLM identyfikuje encje, właściwości i relacje", "Pominięcie, duplikat lub halucynacja"],
        ["4", "Model zwraca dane zgodne z JSON Schema", "Błąd składni albo niedozwolony typ"],
        ["5", "Pydantic i reguły domenowe sprawdzają wynik", "Nieistniejący identyfikator lub zła jednostka"],
        ["6", "Warstwa repozytorium zapisuje graf", "Naruszenie constraint lub błąd transakcji"],
    ], widths=[0.45, 3.4, 3.0], center_cols=(0,))

    doc.add_heading("Wstępny kontrakt JSON", level=2)
    add_code(doc, '{\n  "network_id": "example_01",\n  "nodes": [\n    {"id": "B1", "type": "Bus", "properties": {"vn_kv": 110}},\n    {"id": "L1", "type": "Line", "properties": {"length_km": 10}}\n  ],\n  "relationships": [\n    {"type": "FROM_BUS", "from": "L1", "to": "B1"}\n  ]\n}')
    add_body(doc, "Schemat odpowiedzi powinien ograniczać wartości pola type do uzgodnionego słownika. Walidacja semantyczna musi dodatkowo sprawdzać istnienie wskazanych obiektów, dopuszczalne kierunki relacji, jednostki oraz wartości wymagane dla poszczególnych typów.")

    doc.add_page_break()
    doc.add_heading("Wstępny model bazy grafowej", level=1)
    add_body(doc, "Proponowany model traktuje linie i transformatory jako osobne węzły. Dzięki temu ich parametry techniczne pozostają bezpośrednio dostępne, a odwzorowanie jest zbliżone do struktury tabel pandapower. Rozwiązanie upraszcza również ocenę, czy LLM poprawnie rozpoznał konkretny element infrastruktury.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(schema), width=Inches(6.65))
    add_caption(doc, "Rysunek 2  Wstępny schemat węzłów i relacji Neo4j")

    doc.add_heading("Typy węzłów", level=2)
    add_table(doc, ["Etykieta", "Znaczenie", "Przykładowe właściwości"], [
        ["Bus", "Szyna lub węzeł elektryczny", "source_id, pp_index, name, vn_kv, type, zone, geo, in_service"],
        ["Line", "Linia elektroenergetyczna", "source_id, pp_index, length_km, r_ohm_per_km, x_ohm_per_km, max_i_ka"],
        ["Transformer", "Transformator dwuuzwojeniowy", "source_id, pp_index, sn_mva, vn_hv_kv, vn_lv_kv, vk_percent, tap_pos"],
        ["Generator", "Generator sterowany", "source_id, pp_index, p_mw, vm_pu, min_q_mvar, max_q_mvar"],
        ["ExternalGrid", "Węzeł bilansujący lub sieć zewnętrzna", "source_id, pp_index, vm_pu, va_degree"],
        ["Load", "Odbiór mocy", "source_id, pp_index, p_mw, q_mvar, scaling, in_service"],
        ["Shunt", "Element poprzeczny", "source_id, pp_index, p_mw, q_mvar, step"],
    ], widths=[1.25, 2.0, 3.6])

    doc.add_heading("Typy relacji", level=2)
    add_table(doc, ["Relacja", "Źródło", "Cel", "Znaczenie"], [
        ["FROM_BUS", "Line", "Bus", "Początkowa szyna linii"],
        ["TO_BUS", "Line", "Bus", "Końcowa szyna linii"],
        ["HV_BUS", "Transformer", "Bus", "Strona wysokiego napięcia"],
        ["LV_BUS", "Transformer", "Bus", "Strona niskiego napięcia"],
        ["CONNECTED_TO", "Generator, Load, ExternalGrid, Shunt", "Bus", "Przyłączenie urządzenia do szyny"],
    ], widths=[1.25, 2.1, 1.0, 2.5], center_cols=(0, 2))

    doc.add_page_break()
    doc.add_heading("Ograniczenia i zapis w Neo4j", level=1)
    add_body(doc, "W pierwszej wersji wystarczą ograniczenia unikalności dla source_id w każdej etykiecie. Właściwości wymagane i typy danych można sprawdzać przede wszystkim w warstwie Pydantic, co pozwoli zachować zgodność także w Neo4j Community Edition.")
    add_code(doc, "CREATE CONSTRAINT bus_source_id IF NOT EXISTS\nFOR (n:Bus) REQUIRE n.source_id IS UNIQUE;\n\nCREATE CONSTRAINT line_source_id IF NOT EXISTS\nFOR (n:Line) REQUIRE n.source_id IS UNIQUE;\n\nCREATE CONSTRAINT transformer_source_id IF NOT EXISTS\nFOR (n:Transformer) REQUIRE n.source_id IS UNIQUE;")
    add_body(doc, "Import powinien używać MERGE dla elementów identyfikowanych przez source_id oraz transakcji obejmującej spójny fragment sieci. Dane referencyjne i wyniki LLM warto rozdzielić polem graph_variant albo osobnymi bazami, aby przypadkowo ich nie połączyć.")

    doc.add_heading("Przykładowe wzorce Cypher", level=2)
    add_code(doc, "MERGE (b:Bus {source_id: $source_id})\nSET b += $properties;\n\nMATCH (l:Line {source_id: $line_id})\nMATCH (b:Bus {source_id: $bus_id})\nMERGE (l)-[:FROM_BUS]->(b);")

    doc.add_heading("Otwarte decyzje dotyczące modelu", level=2)
    add_bullets(doc, [
        "Czy zachować osobne etykiety Generator i ExternalGrid, czy wprowadzić wspólną etykietę Source z etykietami dodatkowymi.",
        "Czy wyniki obliczeń pandapower przechowywać jako właściwości, osobne węzły pomiarowe lub poza grafem pierwszej wersji.",
        "Czy relacje mają zawierać metadane pochodzenia, pewność ekstrakcji i fragment tekstu źródłowego.",
        "Czy różne warianty grafu przechowywać w jednej bazie z network_id, czy w osobnych bazach Neo4j.",
    ])

    doc.add_page_break()
    doc.add_heading("Potencjalny widok aplikacji", level=1)
    add_body(doc, "Interfejs pierwszej wersji może składać się z jednego ekranu roboczego. Użytkownik wprowadza opis, wybiera model i uruchamia generowanie. System pokazuje wynik walidacji, graf oraz właściwości wybranego elementu. Dodatkowe funkcje mogą obejmować eksport JSON, porównanie z grafem referencyjnym i listę wykrytych błędów.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ui), width=Inches(6.65))
    add_caption(doc, "Rysunek 3  Robocza makieta głównego widoku aplikacji")
    add_table(doc, ["Obszar widoku", "Zakres funkcji"], [
        ["Panel wejściowy", "Opis tekstowy, wybór modelu, parametry promptu i uruchomienie procesu."],
        ["Podgląd grafu", "Węzły, relacje, filtrowanie typów i zaznaczanie elementów."],
        ["Panel szczegółów", "Właściwości wybranego węzła lub relacji oraz status walidacji."],
        ["Wynik procesu", "Lista ostrzeżeń, błędów, czas odpowiedzi i możliwość eksportu."],
    ], widths=[1.7, 5.15])

    doc.add_page_break()
    doc.add_heading("Plan ewaluacji", level=1)
    add_body(doc, "Ewaluacja powinna oddzielać poprawność struktury od poprawności właściwości. Sam fakt utworzenia prawidłowego JSON nie oznacza, że graf wiernie odtwarza opis. Dla każdego testu należy zachować tekst wejściowy, oczekiwany graf, odpowiedź modelu, wynik walidacji i listę różnic.")
    add_table(doc, ["Poziom", "Metryka", "Interpretacja"], [
        ["Węzły", "precision, recall, F1", "Czy model wykrywa właściwe elementy i nie dodaje nieistniejących."],
        ["Relacje", "precision, recall, F1", "Czy połączenia, kierunki i typy relacji są poprawne."],
        ["Właściwości", "exact match lub błąd liczbowy", "Czy wartości i jednostki odpowiadają opisowi."],
        ["Walidacja", "odsetek odpowiedzi poprawnych za pierwszym razem", "Jak często wynik przechodzi schemat bez naprawy."],
        ["Stabilność", "zgodność między powtórzeniami", "Jak bardzo wynik zależy od losowości modelu."],
        ["Efektywność", "czas, liczba tokenów, koszt", "Jaki jest koszt uzyskania pojedynczego grafu."],
    ], widths=[1.15, 2.15, 3.55])

    doc.add_heading("Proponowane warianty eksperymentów", level=2)
    add_bullets(doc, [
        "Opis prosty, w którym każde zdanie przedstawia jeden obiekt albo relację.",
        "Opis z parametrami technicznymi, jednostkami i kilkoma obiektami w jednym zdaniu.",
        "Opis niepełny lub niejednoznaczny, wymagający zgłoszenia braków zamiast zgadywania.",
        "Porównanie zero shot, few shot oraz structured outputs.",
        "Porównanie co najmniej jednego modelu lokalnego i jednego modelu dostępnego przez API.",
        "Powtórzenie tego samego testu dla kilku ustawień temperatury i kilku losowych uruchomień.",
    ])

    doc.add_heading("Kategorie błędów", level=2)
    add_table(doc, ["Kod", "Kategoria", "Przykład"], [
        ["E1", "Pominięcie obiektu", "Brak transformatora opisanego w tekście"],
        ["E2", "Obiekt nadmiarowy", "Generator niewystępujący w opisie"],
        ["E3", "Błędny typ", "Linia rozpoznana jako transformator"],
        ["E4", "Błędna relacja", "Nieprawidłowa szyna początkowa"],
        ["E5", "Błędna właściwość", "Niepoprawna długość lub moc"],
        ["E6", "Błąd jednostki", "MW zinterpretowane jako kW"],
        ["E7", "Błąd formatu", "Odpowiedź niezgodna z JSON Schema"],
    ], widths=[0.55, 2.15, 4.15], center_cols=(0,))

    doc.add_page_break()
    doc.add_heading("Plan dalszych prac", level=1)
    add_table(doc, ["Etap", "Zakres", "Rezultat"], [
        ["1 Analiza danych", "Pełny wydruk kolumn case14, wybór pól i przypięcie wersji pandapower", "Specyfikacja mapowania"],
        ["2 Model Neo4j", "Etykiety, relacje, source_id, constraints i reguły walidacji", "Zatwierdzony schemat"],
        ["3 Graf referencyjny", "Konwerter DataFrame do Neo4j oraz testy integralności", "Referencyjna sieć case14"],
        ["4 Integracja LLM", "Adapter modeli, prompt, JSON Schema i modele Pydantic", "Pierwsza ekstrakcja z tekstu"],
        ["5 Aplikacja", "API, ekran roboczy, wizualizacja i eksport", "Prototyp demonstracyjny"],
        ["6 Ewaluacja", "Zbiór opisów, porównanie modeli i analiza błędów", "Wyniki do rozdziału badawczego"],
        ["7 Rozszerzenie", "Dodatkowe sieci, dane autentyczne lub fine tuning", "Opcjonalny zakres końcowy"],
    ], widths=[1.2, 3.75, 1.9])

    doc.add_heading("Najbliższe zadania", level=2)
    add_bullets(doc, [
        "Zapisać pełne listy kolumn dla tabel case14 i określić minimalne właściwości każdego typu.",
        "Zatwierdzić reprezentowanie linii i transformatorów jako osobnych węzłów.",
        "Przygotować klasy Pydantic wspólne dla importu referencyjnego i wyniku LLM.",
        "Uruchomić Neo4j lokalnie i utworzyć constraints.",
        "Zaimplementować import szyn, a następnie kolejnych typów elementów wraz z relacjami.",
        "Przygotować pierwszy opis tekstowy odtwarzający wybrany fragment case14.",
    ])

    doc.add_heading("Pytania do promotora", level=2)
    add_bullets(doc, [
        "Czy przedstawiony zakres prototypu jest wystarczający dla pracy inżynierskiej?",
        "Czy linie i transformatory powinny pozostać osobnymi węzłami grafu?",
        "Które parametry techniczne muszą znaleźć się w minimalnym schemacie pierwszej wersji?",
        "Czy case14 wystarczy do demonstracji, jeśli ewaluacja obejmie dodatkowe sieci testowe?",
        "Czy prototyp ma wykonywać obliczenia rozpływu mocy, czy skupić się na ekstrakcji, grafie i analizie topologii?",
        "Ile modeli oraz strategii promptowania należy porównać, aby część eksperymentalna była wystarczająca?",
    ])

    doc.add_page_break()
    doc.add_heading("Podsumowanie", level=1)
    add_body(doc, "Koncepcja systemu jest już określona na poziomie przepływu danych i podziału odpowiedzialności. Najważniejszym elementem pierwszej implementacji będzie wspólny model domenowy używany zarówno przez konwerter pandapower, jak i ścieżkę LLM. Pozwoli on porównywać wyniki bez dodatkowych transformacji oraz precyzyjnie klasyfikować błędy.")
    add_body(doc, "Rozpoczęcie od case14 ogranicza złożoność i zapewnia znany punkt odniesienia. Po zbudowaniu grafu referencyjnego kolejne prace mogą koncentrować się na jakości ekstrakcji, niezawodności walidacji i porównaniu modeli. Dodatkowe sieci i dane rzeczywiste powinny zostać włączone dopiero po ustabilizowaniu schematu oraz procesu ewaluacji.")

    doc.add_heading("Źródła", level=1)
    refs = [
        "Dokumentacja pandapower, Power System Test Cases, case14, https://pandapower.readthedocs.io/en/v3.5.6/networks/power_system_test_cases.html",
        "Repozytorium pandapower, https://github.com/e2nIEE/pandapower",
        "Neo4j, Graph database concepts, https://neo4j.com/docs/getting-started/appendix/graphdb-concepts/",
        "Neo4j Cypher Manual, Create constraints, https://neo4j.com/docs/cypher-manual/current/schema/constraints/create-constraints/",
        "Materiały robocze i wyniki analizy pandapower case14 przekazane w ramach przygotowania pracy.",
    ]
    for idx, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(f"{idx}. {ref}")
        set_run_font(r, size=9.5)

    # Core document metadata without personal details.
    props = doc.core_properties
    props.title = "Raport początkowy pracy inżynierskiej"
    props.subject = "System generowania grafów infrastruktury energetycznej z wykorzystaniem modeli LLM"
    props.keywords = "LLM, pandapower, Neo4j, infrastruktura energetyczna, grafy"

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()
